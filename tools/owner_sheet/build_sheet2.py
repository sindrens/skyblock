import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
qs = json.load(open('final_q.json')); OUT = sys.argv[1]
F='Arial'; HDR=PatternFill('solid', fgColor='C8102E'); HF=Font(name=F,bold=True,color='FFFFFF'); IN=PatternFill('solid', fgColor='FFF7CC')
thin=Side(style='thin', color='D0D0D0'); B=Border(left=thin,right=thin,top=thin,bottom=thin); W=Alignment(wrap_text=True, vertical='top')
wb=Workbook(); ws=wb.active; ws.title='How to answer'
txt=[('Phoenix Skyblock · decisions', Font(name=F,bold=True,size=16,color='C8102E')),
 ('Each row is a choice the game uses right now. Answer in the yellow column with one pick:',None),
 ('   Suggested = keep it as written in "Suggested".',None),('   No = don\'t do it that way.',None),
 ('   Other = write what you want in "Note".',None),('Empty rows keep the suggested answer. Filter Priority = "before launch" to do the important ones first.',None),
 ('Faster on a phone: the clickable page "Phoenix Skyblock Decisions" has the same questions with one tap per answer.',None),('',None),
 ('Example of a filled row:',Font(name=F,bold=True))]
for i,(t,f) in enumerate(txt,1): ws.cell(row=i,column=1,value=t).font=f or Font(name=F,size=11)
for j,(h,v) in enumerate(zip(['Question','Suggested','Your answer','Note'],['Keep the AFK pool limit at 8 hours per day?','Yes, 8 hours','Other','6 hours']),1):
    a=ws.cell(row=10,column=j,value=h); a.font=HF; a.fill=HDR; a.border=B
    b=ws.cell(row=11,column=j,value=v); b.font=Font(name=F); b.border=B
    if j>=3: b.fill=IN
for c,w in zip('ABCD',[60,22,14,24]): ws.column_dimensions[c].width=w
q=wb.create_sheet('Questions')
heads=['#','Area','Feature','Question','Suggested','Your answer','Note','Priority','Context']
for j,(h,w) in enumerate(zip(heads,[7,18,24,62,30,14,30,13,40]),1):
    c=q.cell(row=1,column=j,value=h); c.font=HF; c.fill=HDR; c.border=B; q.column_dimensions[get_column_letter(j)].width=w
for i,x in enumerate(qs,2):
    vals=[x['id'],x['a'],x['f'],x['q'],x['s'],None,None,'before launch' if x['p'] else 'can wait',x['c']]
    for j,v in enumerate(vals,1):
        c=q.cell(row=i,column=j,value=v); c.font=Font(name=F,size=10,bold=(j==4)); c.alignment=W; c.border=B
        if j in (6,7): c.fill=IN
    if x['p']: q.cell(row=i,column=8).font=Font(name=F,size=10,bold=True,color='C8102E')
n=len(qs)+1
dv=DataValidation(type='list', formula1='"Suggested,No,Other"', allow_blank=True); q.add_data_validation(dv); dv.add(f'F2:F{n}')
q.conditional_formatting.add(f'A2:I{n}', FormulaRule(formula=['$F2="Suggested"'], fill=PatternFill('solid', fgColor='DFF5E1')))
q.conditional_formatting.add(f'A2:I{n}', FormulaRule(formula=['$F2="No"'], fill=PatternFill('solid', fgColor='E9EBEF')))
q.conditional_formatting.add(f'A2:I{n}', FormulaRule(formula=['$F2="Other"'], fill=PatternFill('solid', fgColor='FBF0DC')))
q.freeze_panes='E2'; q.auto_filter.ref=f'A1:I{n}'
s=wb.create_sheet('Summary'); wb.move_sheet('Summary', offset=-1)
for j,(h,w) in enumerate(zip(['Area','Questions','Answered','Answered %','Before launch','Before launch answered'],[26,12,12,13,15,22]),1):
    c=s.cell(row=1,column=j,value=h); c.font=HF; c.fill=HDR; c.border=B; s.column_dimensions[get_column_letter(j)].width=w
areas=list(dict.fromkeys(x['a'] for x in qs))
for i,a in enumerate(areas,2):
    s.cell(row=i,column=1,value=a)
    s.cell(row=i,column=2,value=f'=COUNTIF(Questions!$B$2:$B${n},A{i})')
    s.cell(row=i,column=3,value=f'=COUNTIFS(Questions!$B$2:$B${n},A{i},Questions!$F$2:$F${n},"<>")')
    s.cell(row=i,column=4,value=f'=IF(B{i}=0,0,C{i}/B{i})')
    s.cell(row=i,column=5,value=f'=COUNTIFS(Questions!$B$2:$B${n},A{i},Questions!$H$2:$H${n},"before launch")')
    s.cell(row=i,column=6,value=f'=COUNTIFS(Questions!$B$2:$B${n},A{i},Questions!$H$2:$H${n},"before launch",Questions!$F$2:$F${n},"<>")')
t=len(areas)+2; s.cell(row=t,column=1,value='Total')
for col in 'BCEF': s[f'{col}{t}']=f'=SUM({col}2:{col}{t-1})'
s[f'D{t}']=f'=IF(B{t}=0,0,C{t}/B{t})'
for row in s.iter_rows(min_row=2,max_row=t):
    for c in row:
        c.border=B; c.font=Font(name=F,bold=(c.row==t))
        if c.column==4: c.number_format='0%'
wb.save(OUT); print(n-1)
