import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

D = '/tmp/claude-0/-home-claude/e0e0d16a-5da7-5cc0-8715-921df88d8d9f/scratchpad/sheet/'
REPO = '/home/claude/skyblock/'
OUT = sys.argv[1]

# ---------- areas (coarse groups) ----------
AREAS = [
    ('Look & feel', ['look', 'v2', 'colors', 'style', 'resource pack', 'scoreboard', 'tab list', 'motd', 'menu', '/help', 'settings']),
    ('Profiles', ['profile']),
    ('Islands', ['island', 'visiting', 'warps', 'notes', 'disband']),
    ('Social & chat', ['social', 'chat', 'private messages']),
    ('Teleports & essentials', ['teleport', 'essentials', '/spawn']),
    ('Ranks, perks & cosmetics', ['rank', 'perk', 'titles', 'tags', 'achievement', 'cosmetic']),
    ('Worlds & builds', ['world', 'build', 'spawn']),
    ('Staff & punishments', ['staff', 'punish', 'logs', 'appeal', 'freeze']),
    ('Economy', ['economy', 'balance', '/pay', 'money', 'bank', 'interest', 'sell', 'shop', 'auction', 'trad', 'dashboard', 'price']),
    ('The Depths', ['depths', 'satchel']),
    ('Tools, skills & enchants', ['forge', 'tool', 'blueprint', 'skin', 'skill', 'enchant']),
    ('Wardrobe & sets', ['wardrobe', 'armor set', 'loadout']),
    ('Stacking & spawners', ['stack', 'spawner']),
    ('AFK pool', ['afk']),
]
OLD2AREA = {
    'Look & feel (colors, lores, menus)': 'Look & feel', 'Settings': 'Look & feel', 'Permissions & ranks': 'Ranks, perks & cosmetics',
    'Achievements, titles & cosmetics': 'Ranks, perks & cosmetics', 'Teleports & warps': 'Teleports & essentials', 'Profiles': 'Profiles',
    'Islands': 'Islands', 'Island upgrades & stats': 'Islands', 'Economy: money, pay, notes, limits': 'Economy', 'Bank': 'Economy',
    'Sell': 'Economy', 'Shop': 'Economy', 'Auction house': 'Economy', 'Trading': 'Economy', 'Economy dashboard & alerts': 'Economy',
    'The Depths': 'The Depths', 'Satchel': 'The Depths', 'Forge & tools': 'Tools, skills & enchants', 'Skills': 'Tools, skills & enchants',
    'Enchants': 'Tools, skills & enchants', 'Wardrobe & sets': 'Wardrobe & sets', 'Stacking & spawners': 'Stacking & spawners',
    'Punishments & staff': 'Staff & punishments', 'AFK pool': 'AFK pool', 'Leaderboards': 'Look & feel', 'Network & website': 'Other', 'Other': 'Other'}

MATCH_ORDER = ['Profiles', 'Islands', 'AFK pool', 'Economy', 'The Depths', 'Tools, skills & enchants', 'Wardrobe & sets',
               'Stacking & spawners', 'Staff & punishments', 'Ranks, perks & cosmetics', 'Social & chat', 'Teleports & essentials',
               'Worlds & builds', 'Look & feel']
FIX = {'Social: private messages, ignore, AFK and info commands': 'Social & chat', 'Chat: format, filters, mentions and staff chat': 'Social & chat',
       'Tab list and nametags': 'Look & feel', 'Staff essentials (teleport, game mode, fly, vanish, inventory, items)': 'Teleports & essentials',
       'Rank perks: /craft, /ec, /hat, /nick, /trash and staff tools': 'Ranks, perks & cosmetics', 'Ranks, the free Ember rank and staff ranks': 'Ranks, perks & cosmetics',
       'Money notes (/withdraw)': 'Economy', 'Bank (personal and island)': 'Economy', 'Worlds, spawn and builds': 'Worlds & builds'}
def area_for(name):
    if name in FIX: return FIX[name]
    low = name.lower()
    akeys = dict(AREAS)
    for a in MATCH_ORDER:
        keys = akeys[a]
        if any(k in low for k in keys):
            return a
    return 'Other'

rows = []
# ---------- 1. existing open questions (built with a default) ----------
old = json.load(open(D + 'q.json'))
for r in old:
    q = r['question']
    if q.startswith('RESOURCE PACK'):
        continue  # split below from the plan's table
    if q.startswith('V2 colors'):
        for qq, sg in [('Sell price lines in red like your examples?', 'Yes, red'),
                       ('Action lines in bold yellow "» RIGHT-CLICK buy 64" (instead of green "CLICK HERE")?', 'Yes, bold yellow'),
                       ('Thin red | bar on every line inside a section (not only on the header)?', 'Yes, every line')]:
            rows.append(dict(area='Look & feel', feature='Colors and lores (V2)', question=qq, suggested=sg, priority='before launch',
                             why='How every item tooltip looks', source='Built with this default (V2)'))
        continue
    rows.append(dict(area=OLD2AREA.get(r['feature'], 'Other'), feature=r['feature'], question=q, suggested=r['suggested'], priority='',
                     why='', source='Built with this default · ' + r['code']))
# resource pack table
txt = open(REPO + 'docs/plans/RESOURCE_PACK.md').read()
for m in re.finditer(r'^\| (RP-D\d+) \| ([^|]+) \| ([^|]+) \|', txt, re.M):
    rows.append(dict(area='Look & feel', feature='Resource pack (plan)', question=m.group(2).strip(), suggested=m.group(3).strip(),
                     priority='can wait', why='Resource pack, after launch of the vanilla look', source='Plan default · ' + m.group(1)))
# ---------- 2. new questions from the feature review ----------
for f in ['review_A.json', 'review_B.json']:
    for feat in json.load(open(D + f))['features']:
        for q in feat['questions']:
            rows.append(dict(area=area_for(feat['feature']), feature=feat['feature'], question=q['question'].strip(),
                             suggested=q.get('suggested', '').strip(), priority=q.get('priority', '').strip(),
                             why=q.get('why', '').strip(), source='New (feature review)'))
# dedupe
seen, uniq = set(), []
for r in rows:
    k = re.sub(r'\W+', ' ', r['question'].lower())[:100]
    if k in seen: continue
    seen.add(k); uniq.append(r)
rows = uniq
AREA_ORDER = [a for a, _ in AREAS] + ['Other']
PRI = {'before launch': 0, 'can wait': 2, '': 1}
rows.sort(key=lambda r: (AREA_ORDER.index(r['area']), r['feature'], PRI.get(r['priority'], 1)))

# ---------- workbook ----------
F = 'Arial'
HDR = PatternFill('solid', fgColor='C8102E'); HFONT = Font(name=F, bold=True, color='FFFFFF', size=11)
INPUT = PatternFill('solid', fgColor='FFF7CC'); thin = Side(style='thin', color='D0D0D0')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')

wb = Workbook()
# Sheet 1: how to answer
ws = wb.active; ws.title = 'How to answer'
lines = [
    ('Phoenix Skyblock · open questions', Font(name=F, bold=True, size=16, color='C8102E')),
    ('Every row is a choice I made without your answer. The game works with the "Suggested / now live" value until you change it.', None),
    ('', None),
    ('How to answer (only the yellow columns):', Font(name=F, bold=True, size=12)),
    ('Your answer: pick Yes, No, Suggested or Other from the dropdown.', None),
    ('   Yes = yes to the question as written.   No = no to the question.', None),
    ('   Suggested = keep what is in the "Suggested / now live" column.', None),
    ('   Other = something else: write it in "Other / your note".', None),
    ('Other / your note: free text. Numbers, names, "ask me later", anything.', None),
    ('Skip what you don\'t care about. Unanswered rows keep the suggested value.', None),
    ('Start with Priority = "before launch" (filter the Priority column). The Summary tab counts what you have answered.', None),
    ('', None),
    ('Columns', Font(name=F, bold=True, size=12)),
    ('Area / Feature: where in the game. Question: the choice. Suggested / now live: what is built today or what I recommend.', None),
    ('Priority: "before launch" = must be decided before players come; "can wait" = later is fine; empty = not rated (older questions).', None),
    ('Source: "Built with this default" = already in the game with that value · "New (feature review)" = found tonight while going through every feature.', None),
    ('', None),
    ('Example (how a filled row looks):', Font(name=F, bold=True, size=12)),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = f or Font(name=F, size=11); c.alignment = Alignment(wrap_text=False)
ex_hdr = ['Question', 'Suggested / now live', 'Your answer', 'Other / your note']
ex = ['Should /pay take a small fee that leaves the economy?', 'No fee, as built', 'Other', '1% fee, only above $1M']
r0 = len(lines) + 1
for j, (h, v) in enumerate(zip(ex_hdr, ex), 1):
    a = ws.cell(row=r0, column=j, value=h); a.font = HFONT; a.fill = HDR; a.border = BORDER
    b = ws.cell(row=r0 + 1, column=j, value=v); b.font = Font(name=F, size=11); b.border = BORDER; b.alignment = wrap
    if j >= 3: b.fill = INPUT
ws.column_dimensions['A'].width = 60; ws.column_dimensions['B'].width = 30; ws.column_dimensions['C'].width = 16; ws.column_dimensions['D'].width = 34

# Sheet 2: questions
qs = wb.create_sheet('Questions')
heads = ['#', 'Area', 'Feature', 'Question', 'Suggested / now live', 'Your answer', 'Other / your note', 'Priority', 'Why it matters', 'Source']
widths = [6, 18, 26, 70, 40, 14, 40, 14, 36, 26]
for j, (h, w) in enumerate(zip(heads, widths), 1):
    c = qs.cell(row=1, column=j, value=h); c.font = HFONT; c.fill = HDR; c.alignment = Alignment(vertical='center', wrap_text=True); c.border = BORDER
    qs.column_dimensions[get_column_letter(j)].width = w
for i, r in enumerate(rows, 2):
    vals = [i - 1, r['area'], r['feature'], r['question'], r['suggested'], None, None, r['priority'], r['why'], r['source']]
    for j, v in enumerate(vals, 1):
        c = qs.cell(row=i, column=j, value=v); c.font = Font(name=F, size=10); c.alignment = wrap; c.border = BORDER
        if j in (6, 7): c.fill = INPUT
    if r['priority'] == 'before launch':
        qs.cell(row=i, column=8).font = Font(name=F, size=10, bold=True, color='C8102E')
n = len(rows) + 1
dv = DataValidation(type='list', formula1='"Yes,No,Suggested,Other"', allow_blank=True, showDropDown=False)
dv.error = 'Pick Yes, No, Suggested or Other'; dv.prompt = 'Yes / No / Suggested / Other'
qs.add_data_validation(dv); dv.add(f'F2:F{n}')
qs.conditional_formatting.add(f'A2:J{n}', FormulaRule(formula=['LEN($F2)>0'], fill=PatternFill('solid', fgColor='DFF5E1')))
qs.freeze_panes = 'E2'; qs.auto_filter.ref = f'A1:J{n}'
qs.row_dimensions[1].height = 30

# Sheet 3: summary (formulas)
sm = wb.create_sheet('Summary')
sh = ['Area', 'Questions', 'Answered', 'Answered %', 'Before launch', 'Before launch answered']
for j, h in enumerate(sh, 1):
    c = sm.cell(row=1, column=j, value=h); c.font = HFONT; c.fill = HDR; c.border = BORDER
    sm.column_dimensions[get_column_letter(j)].width = [28, 12, 12, 13, 15, 22][j - 1]
areas = [a for a in AREA_ORDER if any(r['area'] == a for r in rows)]
for i, a in enumerate(areas, 2):
    sm.cell(row=i, column=1, value=a)
    sm.cell(row=i, column=2, value=f'=COUNTIF(Questions!$B$2:$B${n},A{i})')
    sm.cell(row=i, column=3, value=f'=COUNTIFS(Questions!$B$2:$B${n},A{i},Questions!$F$2:$F${n},"<>")')
    sm.cell(row=i, column=4, value=f'=IF(B{i}=0,0,C{i}/B{i})')
    sm.cell(row=i, column=5, value=f'=COUNTIFS(Questions!$B$2:$B${n},A{i},Questions!$H$2:$H${n},"before launch")')
    sm.cell(row=i, column=6, value=f'=COUNTIFS(Questions!$B$2:$B${n},A{i},Questions!$H$2:$H${n},"before launch",Questions!$F$2:$F${n},"<>")')
t = len(areas) + 2
sm.cell(row=t, column=1, value='Total').font = Font(name=F, bold=True)
for col in 'BCEF':
    sm[f'{col}{t}'] = f'=SUM({col}2:{col}{t - 1})'
sm[f'D{t}'] = f'=IF(B{t}=0,0,C{t}/B{t})'
for row in sm.iter_rows(min_row=2, max_row=t):
    for c in row:
        c.border = BORDER; c.font = Font(name=F, size=11, bold=(c.row == t))
        if c.column == 4: c.number_format = '0%'
wb.move_sheet('Summary', offset=-1)
wb.save(OUT)
import collections
print(len(rows), collections.Counter(r['area'] for r in rows).most_common(), collections.Counter(r['priority'] for r in rows))
